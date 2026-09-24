import docx
from typing import Dict, Any
from backend.app.ai.documents.base import DocumentParser

class DOCXParser(DocumentParser):
    def extract_text(self, file_path: str) -> Dict[str, Any]:
        try:
            doc = docx.Document(file_path)
            full_text = []

            # Extract paragraphs and headings
            for p in doc.paragraphs:
                if p.text.strip():
                    full_text.append(p.text.strip())

            # Extract tables
            for table in doc.tables:
                for row in table.rows:
                    row_data = [cell.text.strip() for cell in row.cells if cell.text.strip()]
                    if row_data:
                        full_text.append(" | ".join(row_data))

            text_content = "\n\n".join(full_text).strip()
            return {
                "raw_text": text_content,
                "needs_ocr": False,
                "metadata": {"paragraphs": len(doc.paragraphs), "tables": len(doc.tables)},
                "error": "Document contains no readable text" if not text_content else None
            }
        except Exception as e:
            return {
                "raw_text": "",
                "needs_ocr": False,
                "metadata": {},
                "error": f"Failed to extract DOCX: {str(e)}"
            }
